from pathlib import Path
import json
import shutil

import pytest

from tools.runtime.template_adapter import InputBlocked, execute_tabular_template


ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "tests" / "fixtures" / "v09_b_regression.csv"


def test_regression_template_executes_with_explicit_binding(tmp_path):
    result = execute_tabular_template(
        ROOT,
        tmp_path,
        tmp_path / "run",
        "linear_regression",
        {"data_path": str(FIXTURE), "target": "target", "features": ["x1", "x2"], "seed": 7},
    )
    assert result.outputs[0]["value"] == "completed"
    assert result.artifacts
    assert any("baseline_predictions.csv" in x["path"] for x in result.artifacts)


def test_missing_target_is_blocked(tmp_path):
    with pytest.raises(InputBlocked, match="target is required"):
        execute_tabular_template(
            ROOT,
            tmp_path,
            tmp_path / "run",
            "linear_regression",
            {"data_path": str(FIXTURE), "features": ["x1", "x2"]},
        )


def test_multiclass_classification_runs_through_adapter(tmp_path):
    """sklearn >= 1.7 requires an explicit multi_class strategy for multiclass
    roc_auc; the plain "roc_auc" scorer raised
    ValueError: multi_class must be in ('ovo', 'ovr'). Any target with more than
    two classes — including the five keyword classes a D/E problem asks for —
    crashed the classification template."""
    import random

    rng = random.Random(7)
    source = tmp_path / "keywords.csv"
    lines = ["spend,clicks,bucket"]
    for _ in range(80):
        spend = round(rng.uniform(0.0, 100.0), 2)
        clicks = rng.randint(0, 40)
        bucket = min(int(spend // 20), 4)
        lines.append(f"{spend},{clicks},{bucket}")
    source.write_text("\n".join(lines) + "\n", encoding="utf-8")

    result = execute_tabular_template(
        ROOT,
        tmp_path,
        tmp_path / "run",
        "logistic_classification",
        {"data_path": str(source), "target": "bucket", "features": ["spend", "clicks"], "seed": 7},
    )

    assert result.outputs[0]["value"] == "completed"
    assert any("classification_cv" in x["path"] for x in result.artifacts)


def _make_workbook(path):
    """Two sheets, with the trailing-space header the real 2026E Sheet2 has."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    first = workbook.create_sheet("Sheet1")
    first.append(["方案ID", "消费额"])
    for i in range(5):
        first.append([i, i * 10])
    second = workbook.create_sheet("Sheet2")
    second.append(["日期", "新注册数 "])
    for i in range(40):
        second.append([f"2025-01-{i % 28 + 1:02d}", 100 + (i * 7) % 50])
    workbook.save(path)


def test_excel_binding_can_target_a_named_sheet(tmp_path):
    """_prepare_csv used to hard-code sheet_name=0, so Sheet2/Sheet3 of a real
    CUMCM attachment were unreachable from the compute stage. The header also
    carries a trailing space there, which DataProfile strips."""
    book = tmp_path / "attach.xlsx"
    _make_workbook(book)

    result = execute_tabular_template(
        ROOT,
        tmp_path,
        tmp_path / "run",
        "time_series_baseline",
        {"data_path": str(book), "sheet": "Sheet2", "time_col": "日期", "target": "新注册数", "seed": 7},
    )

    assert result.outputs[0]["value"] == "completed"
    manifest = json.loads(
        (tmp_path / "run" / "results" / "time_series_manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["parameters"]["time_col"] == "日期"
    assert manifest["parameters"]["features"] == ["lag1_target"]
    assert manifest["parameters"]["derived_features"] == ["lag1_target = target(t - 1)"]


def test_excel_binding_accepts_a_sheet_index(tmp_path):
    book = tmp_path / "attach.xlsx"
    _make_workbook(book)

    # Sheet2 by 0-based index. If the wrong sheet were read the target column
    # would not exist and the binding would be refused.
    result = execute_tabular_template(
        ROOT,
        tmp_path,
        tmp_path / "run",
        "linear_regression",
        {"data_path": str(book), "sheet": 1, "target": "新注册数", "seed": 7},
    )

    assert result.outputs[0]["value"] == "completed"


def test_toy_template_is_refused_instead_of_fabricating_a_result(tmp_path):
    """optimization.py / monte_carlo.py / sensitivity.py hard-code their objective
    and read no external input. Wiring them to a real binding would emit
    fabricated numbers with a valid run manifest and RUN_COMPLETE status."""
    with pytest.raises(InputBlocked, match="would fabricate a result"):
        execute_tabular_template(
            ROOT,
            tmp_path,
            tmp_path / "run",
            "optimization",
            {"data_path": str(FIXTURE), "target": "target"},
        )


def test_clustering_reports_that_no_data_driven_implementation_exists(tmp_path):
    with pytest.raises(InputBlocked, match="no data-driven implementation"):
        execute_tabular_template(
            ROOT,
            tmp_path,
            tmp_path / "run",
            "clustering",
            {"data_path": str(FIXTURE), "target": "target"},
        )


def test_evaluation_binding_ranks_entities_with_entropy_topsis(tmp_path):
    """Evaluation bindings have no dependent variable, so they must not be forced
    through the target/features contract."""
    source = tmp_path / "units.csv"
    lines = ["单元,投入,产出,曝光"]
    for index, (cost, gain, views) in enumerate([(3, 30, 1000), (1, 10, 100), (2, 25, 300)], 1):
        lines.append(f"U{index},{cost},{gain},{views}")
    source.write_text("\n".join(lines) + "\n", encoding="utf-8")

    result = execute_tabular_template(
        ROOT,
        tmp_path,
        tmp_path / "run",
        "entropy_topsis",
        {
            "data_path": str(source),
            "entity": "单元",
            "indicators": [
                {"name": "投入", "direction": "-"},
                {"name": "产出", "direction": "+"},
                {"name": "点击率", "direction": "+", "ratio": ["产出", "曝光"]},
            ],
        },
    )

    assert result.outputs[0]["value"] == "completed"
    manifest = json.loads(
        (tmp_path / "run" / "results" / "entropy_topsis_manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["parameters"]["entity_col"] == "单元"
    assert manifest["parameters"]["entities"] == 3
    assert sum(manifest["parameters"]["weights"].values()) == pytest.approx(1.0)
    assert (tmp_path / "run" / "results" / "entropy_topsis_scores.csv").exists()


def test_evaluation_binding_requires_an_entity_and_known_columns(tmp_path):
    source = tmp_path / "units.csv"
    source.write_text("单元,投入\nU1,3\nU2,1\n", encoding="utf-8")

    with pytest.raises(InputBlocked, match="entity is required"):
        execute_tabular_template(
            ROOT,
            tmp_path,
            tmp_path / "run-a",
            "entropy_topsis",
            {"data_path": str(source), "indicators": [{"name": "投入"}]},
        )

    with pytest.raises(InputBlocked, match="columns not found"):
        execute_tabular_template(
            ROOT,
            tmp_path,
            tmp_path / "run-b",
            "entropy_topsis",
            {"data_path": str(source), "entity": "单元", "indicators": [{"name": "不存在"}]},
        )
