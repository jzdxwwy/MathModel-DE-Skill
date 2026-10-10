"""Deterministic file inspection primitives for V0.7.

The inspector never invents semantic meaning. It is deliberately streaming-aware
so large CUMCM attachments do not get loaded into memory during ingestion.
"""
from __future__ import annotations

import csv
import hashlib
import json
import mimetypes
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

MAX_TEXT_BYTES = 1 * 1024 * 1024
MAX_JSON_PROFILE_BYTES = 20 * 1024 * 1024
DTYPE_SAMPLE_ROWS = 200


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _text_file(path: Path, limit: int = MAX_TEXT_BYTES) -> str:
    with path.open("rb") as f:
        data = f.read(limit + 1)
    text = data[:limit].decode("utf-8", errors="replace")
    if len(data) > limit:
        text += "\n[INGESTION_PREVIEW_TRUNCATED]"
    return text


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = []
    for p in root.findall(".//w:p", ns):
        text = "".join(t.text or "" for t in p.findall(".//w:t", ns))
        if text.strip():
            paragraphs.append(text.strip())
    return "\n".join(paragraphs)


def _csv_profile(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        sample = f.read(64 * 1024)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        header = next(reader, [])
        rows = 0
        empty_cells = 0
        # Exact duplicate detection is intentionally bounded. Large files are
        # profiled without materializing all rows or a million-row hash set.
        seen_sample: set[tuple[str, ...]] = set()
        duplicate_sample = 0
        for row in reader:
            rows += 1
            empty_cells += sum(1 for cell in row if not str(cell).strip())
            if rows <= 5000:
                key = tuple(row)
                if key in seen_sample:
                    duplicate_sample += 1
                seen_sample.add(key)
    return {
        "rows": rows,
        "columns": len(header),
        "columns_profile": [{"name": str(name), "dtype": "string"} for name in header],
        "delimiter": getattr(dialect, "delimiter", ","),
        "empty_cells": empty_cells,
        "duplicate_rows_sample": duplicate_sample,
        "duplicate_scope": "first_5000_rows_only",
    }


def _json_profile(path: Path) -> dict[str, Any]:
    if path.stat().st_size > MAX_JSON_PROFILE_BYTES:
        return {"profile_skipped": True, "reason": "JSON file exceeds safe deterministic profile size"}
    value = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    if isinstance(value, list):
        rows = value
    elif isinstance(value, dict):
        rows = value.get("data") if isinstance(value.get("data"), list) else []
    else:
        rows = []
    if rows and all(isinstance(x, dict) for x in rows):
        names = sorted({k for row in rows for k in row})
        profile = [{"name": k, "dtype": "mixed"} for k in names]
    else:
        profile = []
    return {"rows": len(rows), "columns": len(profile), "columns_profile": profile}


def _infer_dtype(values: list[Any]) -> str:
    """Coarse deterministic dtype from observed cell values. Facts only."""
    observed = {type(v) for v in values if v is not None and str(v).strip() != ""}
    if not observed:
        return "unknown"
    if observed <= {bool}:
        return "boolean"
    if observed <= {int, float, bool}:
        return "number"
    if observed <= {datetime, date}:
        return "datetime"
    if observed <= {str}:
        return "string"
    return "mixed"


def _xlsx_profile_stdlib(path: Path) -> dict[str, Any]:
    """Row/cell counting fallback used when openpyxl is not importable.

    Without openpyxl the shared-string table cannot be resolved, so column
    headers are unavailable. ``profile_limited`` records that fact so the caller
    reports a risk instead of a silent pass.
    """
    with zipfile.ZipFile(path) as z:
        sheet_files = sorted(n for n in z.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n))
        sheets = []
        for index, sheet in enumerate(sheet_files):
            root = ET.fromstring(z.read(sheet))
            xml_rows = len(root.findall(".//{*}sheetData/{*}row"))
            sheets.append({"name": sheet, "index": index, "rows": max(xml_rows - 1, 0), "columns": 0})
    primary = sheets[0] if sheets else {"name": None, "index": 0, "rows": 0, "columns": 0}
    return {
        "rows": primary["rows"],
        "columns": primary["columns"],
        "columns_profile": [],
        "primary_sheet": primary["name"],
        "sheets": sheets,
        "profile_limited": "openpyxl is not installed; column headers could not be read",
    }


def _xlsx_profile(path: Path) -> dict[str, Any]:
    """Profile an .xlsx workbook into the shared tabular contract.

    Returns the same top-level keys as ``_csv_profile``/``_json_profile``
    (``rows``/``columns``/``columns_profile``) so the DataProfile builder can
    actually consume the result, plus a per-sheet breakdown. The top-level shape
    describes the first sheet in workbook order; a workbook with more than one
    sheet is reported through ``sheets`` and must be surfaced as a data risk.

    ``rows`` counts data rows excluding the header, matching ``_csv_profile``.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return _xlsx_profile_stdlib(path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for index, worksheet in enumerate(workbook.worksheets):
            header: list[str] = []
            samples: list[list[Any]] = []
            data_rows = 0
            for values in worksheet.iter_rows(values_only=True):
                if not header:
                    if all(v is None or not str(v).strip() for v in values):
                        continue  # leading blank rows before the header
                    header = [
                        str(v).strip() if v is not None and str(v).strip() else f"column_{i + 1}"
                        for i, v in enumerate(values)
                    ]
                    samples = [[] for _ in header]
                    continue
                data_rows += 1
                if data_rows <= DTYPE_SAMPLE_ROWS:
                    for i in range(min(len(header), len(values))):
                        samples[i].append(values[i])
            sheets.append({
                "name": worksheet.title,
                "index": index,
                "rows": data_rows,
                "columns": len(header),
                "columns_profile": [
                    {"name": name, "dtype": _infer_dtype(sample)}
                    for name, sample in zip(header, samples)
                ],
            })
    finally:
        workbook.close()

    primary = sheets[0] if sheets else {
        "name": None, "index": 0, "rows": 0, "columns": 0, "columns_profile": [],
    }
    return {
        "rows": primary["rows"],
        "columns": primary["columns"],
        "columns_profile": primary["columns_profile"],
        "primary_sheet": primary["name"],
        "sheets": [
            {
                "name": s["name"],
                "index": s["index"],
                "rows": s["rows"],
                "columns": s["columns"],
                "columns_profile": s["columns_profile"],
            }
            for s in sheets
        ],
    }


def extract_text(path: Path) -> tuple[str, str | None]:
    ext = path.suffix.lower()
    if ext in {".txt", ".md", ".markdown", ".rst", ".log"}:
        return _text_file(path), None
    if ext == ".docx":
        try:
            return _docx_text(path), None
        except Exception as exc:
            return "", f"docx extraction failed: {exc}"
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            chunks = []
            total = 0
            truncated = False
            for page in reader.pages:
                text = page.extract_text() or ""
                if total + len(text) > MAX_TEXT_BYTES:
                    chunks.append(text[: max(0, MAX_TEXT_BYTES - total)])
                    truncated = True
                    break
                chunks.append(text)
                total += len(text)
            out = "\n\n".join(chunks)
            if truncated:
                out += "\n[INGESTION_PREVIEW_TRUNCATED]"
            return out, None
        except ImportError:
            return "", "PDF text extraction requires optional dependency pypdf"
        except Exception as exc:
            return "", f"pdf extraction failed: {exc}"
    if ext in {".csv", ".tsv", ".json"}:
        return _text_file(path), None
    if ext == ".xlsx":
        # Profiled as a tabular workbook by _xlsx_profile. There is no separate
        # plain-text extraction, and reporting one would be a false alarm that
        # dilutes the warnings that actually matter.
        return "", None
    return "", "no text extractor for this format"


def inspect_file(path_like: str) -> dict[str, Any]:
    path = Path(path_like).expanduser().resolve()
    item: dict[str, Any] = {
        "path": str(path), "name": path.name, "extension": path.suffix.lower(),
        "mime": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        "exists": path.exists(),
    }
    if not path.exists() or not path.is_file():
        item["status"] = "MISSING"
        item["error"] = "file does not exist or is not a regular file"
        return item
    item.update({"status": "READABLE", "bytes": path.stat().st_size, "sha256": sha256_file(path)})
    text, extraction_error = extract_text(path)
    item["text"] = text
    item["text_chars"] = len(text)
    if extraction_error:
        item["extraction_warning"] = extraction_error
    try:
        ext = path.suffix.lower()
        if ext in {".csv", ".tsv"}:
            item["data_profile"] = _csv_profile(path)
        elif ext == ".json":
            item["data_profile"] = _json_profile(path)
        elif ext == ".xlsx":
            item["data_profile"] = _xlsx_profile(path)
    except Exception as exc:
        item["data_profile_error"] = str(exc)
    return item
